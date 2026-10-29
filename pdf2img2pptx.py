#!/usr/bin/env python3
"""
PDF to Image to PPTX Converter

Renders PDF pages to PNG images and builds a 16:9 PPTX with one full-slide
image per page.
"""

import argparse
import os
import sys
import tempfile
import shutil
from pathlib import Path
from typing import List, Optional

try:
    import fitz  # PyMuPDF
except ImportError:
    print("Error: PyMuPDF is not installed. Run: pip install PyMuPDF", file=sys.stderr)
    sys.exit(1)

try:
    from pptx import Presentation
    from pptx.util import Inches
except ImportError:
    print("Error: python-pptx is not installed. Run: pip install python-pptx", file=sys.stderr)
    sys.exit(1)

try:
    from PIL import Image
except ImportError:
    print("Error: Pillow is not installed. Run: pip install Pillow", file=sys.stderr)
    sys.exit(1)


SUPPORTED_IMAGE_FORMATS = ('.png', '.jpg', '.jpeg', '.bmp', '.gif')


def pdf_to_images(pdf_path: str, output_dir: str, dpi: int = 600) -> List[str]:
    """
    Render each page of a PDF to a PNG file named page_NNN.png.

    Returns:
        Paths of the written images, in page order.

    Raises:
        FileNotFoundError: The PDF does not exist.
        Exception: PyMuPDF failed to open or render the PDF.
    """
    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"PDFファイルが見つかりません: {pdf_path}")

    os.makedirs(output_dir, exist_ok=True)

    print(f"PDFを開いています: {pdf_path}")

    try:
        doc = fitz.open(pdf_path)
        total_pages = len(doc)
        print(f"総ページ数: {total_pages}")

        image_paths = []

        for page_num in range(total_pages):
            page = doc[page_num]

            # PDF coordinates are in points, 72 per inch.
            mat = fitz.Matrix(dpi / 72, dpi / 72)

            pix = page.get_pixmap(matrix=mat)

            # Zero padding keeps the files in page order when img2pptx sorts them by name.
            output_filename = f"page_{page_num + 1:03d}.png"
            output_path = os.path.join(output_dir, output_filename)

            pix.save(output_path)
            image_paths.append(output_path)

            print(f"  ページ {page_num + 1}/{total_pages} を変換しました: {output_filename}")

        doc.close()
        print(f"全ページの変換が完了しました。出力先: {output_dir}")

        return image_paths

    except Exception as e:
        raise Exception(f"PDF処理中にエラーが発生しました: {str(e)}")


def images_to_pptx(image_source: str, output_pptx: str) -> None:
    """
    Build a PPTX with one slide per image from an image file or a directory.

    Images in a directory are used in file name order.

    Raises:
        FileNotFoundError: The file or directory does not exist.
        ValueError: No file with a supported image extension was found.
    """
    if not os.path.exists(image_source):
        raise FileNotFoundError(f"画像ソースが見つかりません: {image_source}")

    image_files = []

    if os.path.isfile(image_source):
        if image_source.lower().endswith(SUPPORTED_IMAGE_FORMATS):
            image_files = [image_source]
    elif os.path.isdir(image_source):
        for filename in sorted(os.listdir(image_source)):
            if filename.lower().endswith(SUPPORTED_IMAGE_FORMATS):
                image_files.append(os.path.join(image_source, filename))

    if not image_files:
        raise ValueError(f"有効な画像ファイルが見つかりません: {image_source}")

    print(f"画像ファイル数: {len(image_files)}")

    prs = Presentation()

    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    # Layout 6 of the default template is "Blank".
    blank_layout = prs.slide_layouts[6]

    for idx, image_path in enumerate(image_files, start=1):
        print(f"  スライド {idx}/{len(image_files)} を作成中: {os.path.basename(image_path)}")

        with Image.open(image_path) as img:
            img_width, img_height = img.size

        img_aspect = img_width / img_height
        slide_aspect = prs.slide_width / prs.slide_height

        # Scale the image to cover the whole slide without distortion and center it.
        # The parts outside the slide stay in the file but are not shown.
        if img_aspect > slide_aspect:
            height = prs.slide_height
            width = int(height * img_aspect)
            left = int((prs.slide_width - width) / 2)
            top = 0
        else:
            width = prs.slide_width
            height = int(width / img_aspect)
            left = 0
            top = int((prs.slide_height - height) / 2)

        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(image_path, left, top, width, height)

    prs.save(output_pptx)
    print(f"PPTXファイルを保存しました: {output_pptx}")


def pdf_to_pptx(pdf_path: str, output_pptx: str, dpi: int = 600,
                keep_images: bool = False, image_dir: Optional[str] = None) -> None:
    """
    Convert a PDF to a PPTX through PNG images.

    The images are written to image_dir if given, otherwise to a temporary
    directory that is deleted at the end unless keep_images is set.
    """
    if image_dir:
        temp_dir = image_dir
        os.makedirs(temp_dir, exist_ok=True)
        is_temp = False
    else:
        temp_dir = tempfile.mkdtemp(prefix="pdf2pptx_")
        is_temp = True

    try:
        print("ステップ1: PDFを画像に変換中...")
        pdf_to_images(pdf_path, temp_dir, dpi)

        print("\nステップ2: 画像からPPTXを生成中...")
        images_to_pptx(temp_dir, output_pptx)

        print("\n変換が完了しました！")

        if keep_images:
            print(f"中間画像は保持されています: {temp_dir}")

    finally:
        if is_temp and not keep_images:
            shutil.rmtree(temp_dir)
            print(f"一時ファイルを削除しました: {temp_dir}")


def main():
    parser = argparse.ArgumentParser(
        description='PDF to Image to PPTX Converter',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
使用例:
  # PDFを画像に変換
  %(prog)s pdf2img input.pdf --output ./images --dpi 600

  # 画像からPPTXを生成
  %(prog)s img2pptx ./images --output output.pptx

  # PDFから直接PPTXを生成
  %(prog)s pdf2pptx input.pdf --output output.pptx --dpi 600
        """
    )

    subparsers = parser.add_subparsers(dest='command', help='実行するコマンド')

    parser_pdf2img = subparsers.add_parser('pdf2img', help='PDFを画像に変換')
    parser_pdf2img.add_argument('input', help='入力PDFファイルパス')
    parser_pdf2img.add_argument('--output', default='./images', help='出力ディレクトリ（デフォルト: ./images）')
    parser_pdf2img.add_argument('--dpi', type=int, default=600, help='解像度（デフォルト: 600）')

    parser_img2pptx = subparsers.add_parser('img2pptx', help='画像からPPTXを生成')
    parser_img2pptx.add_argument('input', help='画像ファイルまたはディレクトリパス')
    parser_img2pptx.add_argument('--output', required=True, help='出力PPTXファイルパス')

    parser_pdf2pptx = subparsers.add_parser('pdf2pptx', help='PDFから直接PPTXを生成')
    parser_pdf2pptx.add_argument('input', help='入力PDFファイルパス')
    parser_pdf2pptx.add_argument('--output', required=True, help='出力PPTXファイルパス')
    parser_pdf2pptx.add_argument('--dpi', type=int, default=600, help='解像度（デフォルト: 600）')
    parser_pdf2pptx.add_argument('--keep-images', action='store_true', help='中間画像を削除せずに保持')
    parser_pdf2pptx.add_argument('--image-dir', help='中間画像の保存先（デフォルト: 一時ディレクトリ）')

    args = parser.parse_args()

    if not args.command:
        parser.print_help()
        sys.exit(1)

    try:
        if args.command == 'pdf2img':
            pdf_to_images(args.input, args.output, args.dpi)

        elif args.command == 'img2pptx':
            images_to_pptx(args.input, args.output)

        elif args.command == 'pdf2pptx':
            pdf_to_pptx(args.input, args.output, args.dpi,
                       args.keep_images, args.image_dir)

    except Exception as e:
        print(f"エラー: {str(e)}", file=sys.stderr)
        sys.exit(1)


if __name__ == '__main__':
    main()
